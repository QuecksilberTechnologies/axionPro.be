from __future__ import annotations

from pathlib import Path
from tempfile import gettempdir

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "AxionPro Device UI API Developer Guide.docx"
TMP = Path(gettempdir()) / "axionpro-device-ui-guide"
TMP.mkdir(exist_ok=True)

NAVY = "17365D"
BLUE = "1F4E78"
PALE_BLUE = "EAF3F8"
PALE_GREY = "F3F6F8"
GRID = "D9E2F3"
DARK = RGBColor(0, 0, 0)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_border(cell, color: str = "D9D9D9") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), color)


def set_cell_padding(cell, top=90, start=110, bottom=90, end=110) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    mar = tc_pr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tc_pr.append(mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def code_json(value: str) -> str:
    return value.strip()


def font(size: int, bold: bool = False):
    try:
        return ImageFont.truetype("arial.ttf", size), bold
    except OSError:
        return ImageFont.load_default(), bold


def flow_image(path: Path, title: str, nodes: list[str], labels: list[str]) -> None:
    width, height = 1600, 480
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    title_font, _ = font(30, True)
    body_font, _ = font(22)
    draw.text((50, 32), title, fill="#000000", font=title_font)
    y, box_w, box_h = 185, 238, 118
    x_positions = [32, 346, 660, 974, 1288]
    for index, node in enumerate(nodes):
        x = x_positions[index]
        draw.rounded_rectangle((x, y, x + box_w, y + box_h), radius=16, fill="#EAF3F8", outline="#1F4E78", width=3)
        words, lines, line = node.split(), [], ""
        for word in words:
            next_line = (line + " " + word).strip()
            if len(next_line) > 19:
                lines.append(line)
                line = word
            else:
                line = next_line
        if line:
            lines.append(line)
        ty = y + 31 - max(0, len(lines) - 2) * 10
        for line in lines:
            bbox = draw.textbbox((0, 0), line, font=body_font)
            draw.text((x + (box_w - (bbox[2] - bbox[0])) / 2, ty), line, fill="#000000", font=body_font)
            ty += 28
        if index < len(nodes) - 1:
            start, end = x + box_w + 8, x_positions[index + 1] - 10
            draw.line((start, y + box_h / 2, end, y + box_h / 2), fill="#1F4E78", width=5)
            draw.polygon([(end, y + box_h / 2), (end - 16, y + box_h / 2 - 10), (end - 16, y + box_h / 2 + 10)], fill="#1F4E78")
            if index < len(labels):
                label = labels[index]
                bbox = draw.textbbox((0, 0), label, font=body_font)
                draw.text(((start + end - (bbox[2] - bbox[0])) / 2, y + box_h / 2 - 44), label, fill="#17365D", font=body_font)
    image.save(path)


DEVICE_FLOW = TMP / "device_command_flow.png"
EMPLOYEE_FLOW = TMP / "employee_flow.png"
CARD_FLOW = TMP / "card_flow.png"
flow_image(DEVICE_FLOW, "Device command path", ["Angular UI", "API and permission", "Business tables", "DeviceCommand queue", "Device acknowledgement"], [])
flow_image(EMPLOYEE_FLOW, "Employee credential path", ["Employee location", "Enrollment", "Credential request", "Queue and device", "Status refresh"], [])
flow_image(CARD_FLOW, "Card lifecycle", ["Host inventory", "Available card", "Tenant bind", "Assigned card", "Unbind or remove"], [])


doc = Document()
section = doc.sections[0]
section.top_margin = Cm(1.65)
section.bottom_margin = Cm(1.45)
section.left_margin = Cm(1.55)
section.right_margin = Cm(1.55)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
styles["Normal"].font.size = Pt(9.5)
styles["Normal"].font.color.rgb = DARK
for style_name, size in (("Title", 22), ("Heading 1", 15), ("Heading 2", 12), ("Heading 3", 10.5)):
    style = styles[style_name]
    style.font.name = "Aptos Display"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = DARK

footer = section.footer
footer_p = footer.paragraphs[0]
footer_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer_p.add_run("AxionPro UI API Developer Guide  |  Page ").font.size = Pt(8)
add_page_number(footer_p)


def para(text: str = "", style=None, bold_prefix: str | None = None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.08
    if bold_prefix and text.startswith(bold_prefix):
        p.add_run(bold_prefix).bold = True
        p.add_run(text[len(bold_prefix):])
    else:
        p.add_run(text)
    return p


def heading(text: str, level: int = 1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(5)
    return p


def bullets(items: list[str]):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        p.add_run(item)


def table(headers: list[str], rows: list[list[str]], widths: list[float] | None = None, size: float = 8.2):
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    t.style = "Table Grid"
    hdr = t.rows[0]
    for i, label in enumerate(headers):
        cell = hdr.cells[i]
        if widths:
            cell.width = Inches(widths[i])
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_shading(cell, NAVY)
        set_cell_border(cell)
        set_cell_padding(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(label)
        r.bold = True
        r.font.size = Pt(size)
        r.font.color.rgb = RGBColor(255, 255, 255)
    for r_index, values in enumerate(rows):
        cells = t.add_row().cells
        for i, value in enumerate(values):
            cell = cells[i]
            if widths:
                cell.width = Inches(widths[i])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_border(cell)
            set_cell_padding(cell)
            if r_index % 2 == 1:
                set_cell_shading(cell, PALE_GREY)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(value)
            run.font.size = Pt(size)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def code_block(text: str):
    t = doc.add_table(rows=1, cols=1)
    t.autofit = False
    cell = t.cell(0, 0)
    set_cell_shading(cell, "F6F8FA")
    set_cell_border(cell, "C8D0D9")
    set_cell_padding(cell, 100, 130, 100, 130)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run(text.strip())
    run.font.name = "Consolas"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
    run.font.size = Pt(7.3)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def api(title: str, method_path: str, why: str, call_when: str, body: str | None, response: str, db: str, notes: list[str] | None = None):
    heading(title, 3)
    table(["Method and endpoint", "Why and when Angular calls it"], [[method_path, f"Why: {why}\nWhen: {call_when}"]], [2.75, 4.1], 8.2)
    if body:
        para("Request body", bold_prefix="Request body")
        code_block(body)
    para("Successful response example", bold_prefix="Successful response example")
    code_block(response)
    para(f"Database and queue effect: {db}", bold_prefix="Database and queue effect: ")
    if notes:
        bullets(notes)


SUCCESS_QUEUE = code_json('''{
  "isSucceeded": true,
  "message": "Settings have been queued for the device.",
  "data": {
    "deviceCommandId": 981,
    "internalTrackingId": "e1d472c8-6650-4ae3-bb76-8b2d3ce7bcae",
    "status": "Queued"
  },
  "errors": []
}''')

SUCCESS_ENROLLMENT = code_json('''{
  "isSucceeded": true,
  "message": "Employee device enrollment created successfully.",
  "data": {
    "id": "ENROLLMENT_OPAQUE_ID",
    "employeeId": "EMPLOYEE_OPAQUE_ID",
    "employeeName": "Ananya Sharma",
    "tenantDeviceId": "TENANT_DEVICE_OPAQUE_ID",
    "deviceCode": "DELHI-GATE-01",
    "tenantLocationId": "LOCATION_OPAQUE_ID",
    "tenantLocationName": "Delhi Office",
    "tenantCardId": null,
    "maskedCardNumber": null,
    "faceDeploymentStatus": 0,
    "cardDeploymentStatus": 0,
    "pinDeploymentStatus": 0,
    "userActivationCommandStatus": null,
    "accessWindows": [{"dayOfWeek": 1, "startLocalTime": "09:00:00", "endLocalTime": "18:00:00", "isActive": true}],
    "isActive": true
  },
  "errors": []
}''')


# Cover
p = doc.add_paragraph(style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.LEFT
p.add_run("AxionPro Device UI API Developer Guide")
para("For Angular UI developers. This guide explains the recently added device provisioning, runtime settings, employee credential, card inventory, dropdown, work-configuration APIs, and the current attendance-ingestion boundary. Each endpoint section states why it exists, exactly when the UI should call it, a request and response example, and the database or device queue impact.")
table(["Audience", "Scope", "Important result"], [["Angular UI developer", "Device configuration, employee credentials, card inventory, work configuration", "The UI never opens a LAN connection to a device. All remote device actions enter the durable DeviceCommand queue."]], [1.55, 2.9, 2.4], 8.7)
para("Example IDs in this document are placeholders. Use opaque IDs returned by the APIs; do not construct, decode, or replace them in Angular. `moduleId` and `operationId` must be sourced from the authenticated user's permitted module operations.")

heading("How a device command reaches the physical device", 1)
doc.add_picture(str(DEVICE_FLOW), width=Inches(6.85))
para("For HTTPS devices, the device takes the next command on its normal outbound gateway heartbeat. For MQTTS devices, the dispatcher publishes through the broker. The `DeviceCommand` record is created first, and a device acknowledgement creates a `DeviceCommandResponse` record. UI acceptance is not device success; show the command status until it is `Completed` or `Failed`.")

heading("Common request and response rules", 1)
table(["Rule", "UI implementation"], [
    ["Authorization", "Send the normal Bearer token. All secured calls are also checked against the submitted `moduleId` and `operationId`."],
    ["Tenant scope", "Host requests include opaque `tenantId`. A logged-in Tenant admin derives scope from the token and must not select another Tenant."],
    ["Opaque identifiers", "Send opaque employee, device, enrollment and card IDs exactly as returned. Never send the decrypted Employee database ID to a device API."],
    ["Secrets", "`currentWebServerPassword`, a new Web password, a screen PIN, employee PIN, app token, card number and photo must never be displayed from a response or retained in browser state longer than necessary."],
    ["Status refresh", "After any queue-producing endpoint, refresh the related read endpoint. Treat `Queued`, `Publishing`, `AwaitingResponse`, and `RetryScheduled` as in-progress."],
], [1.55, 5.3], 8.5)

heading("Endpoint catalogue", 1)
table(["Area", "Endpoint family", "Primary table or queue"], [
    ["Initial provisioning", "TenantDeviceConfiguration issue-bootstrap-url", "DeviceInitialProvisioning"],
    ["Tenant device runtime", "Gateway, runtime, reboot, location, MQTTS dispatch", "TenantDeviceConfiguration, TenantDevice, DeviceCommand"],
    ["Typed device settings", "time, bell, device setup, advanced, lock, network, local web access", "DeviceCommand only"],
    ["Dropdowns", "device-ddl-options sections", "No database write"],
    ["Host card inventory", "TenantCardMaster CRUD", "TenantCardMaster"],
    ["Employee device credential", "EmployeeDeviceEnrollment and face/card/PIN actions", "EmployeeDeviceEnrollment, EmployeeDeviceAccessWindow, TenantCardMaster, DeviceCommand"],
    ["Employee eligibility", "EmployeeLocationAssignment", "EmployeeLocationAssignment"],
    ["Work configuration", "EmployeeWorkArrangement, EmployeeWorkPattern, EmployeeWorkModeOverride", "Corresponding work tables"],
    ["Attendance punch validation", "Not yet exposed as a production endpoint", "Required future raw-event and attendance-decision records"],
], [1.35, 3.15, 2.35], 8.3)

heading("1 Initial device provisioning and tenant device connection", 1)
api(
    "Issue initial HTTPS bootstrap URL",
    "POST /api/TenantDeviceConfiguration/issue-bootstrap-url",
    "Creates a short-lived initial device gateway address for a brand-new, active, unassigned HTTPS-capable DeviceMaster.",
    "Host admin has physically unpacked a device and must give the technician one address to paste into the device server setting.",
    '''{
  "deviceMasterId": 1,
  "lifetimeMinutes": 120,
  "moduleId": 49,
  "operationId": 4
}''',
    '''{
  "isSucceeded": true,
  "message": "Initial device gateway URL generated. It will not be shown again.",
  "data": {
    "deviceSerialNumber": "AYUC24030780",
    "initialGatewayUrl": "https://axionpro-api.onrender.com/api/initial/AYUC24030780/ONE_TIME_TOKEN",
    "heartbeatIntervalSeconds": 20,
    "expiresDateTime": "2026-09-08T15:30:00Z"
  },
  "errors": []
}''',
    "Inserts one `DeviceInitialProvisioning` record with only the token hash, expiry, issuer and heartbeat. Any still-valid previous bootstrap row for that DeviceMaster is revoked. The raw URL is not persisted and is returned once.",
    ["Do not put this URL in an Angular list, detail page, or audit response.", "This is only for an unassigned physical device. It is not the later Tenant device URL rotation flow."])

api(
    "Read current safe gateway address",
    "GET /api/TenantDeviceConfiguration/gateway-address/{tenantDeviceId}?moduleId={moduleId}&operationId={operationId}",
    "Shows the non-secret server host, path, port and whether a gateway exists or is being replaced.",
    "Tenant device connection screen opens or refreshes.",
    None,
    '''{
  "isSucceeded": true,
  "message": "Device gateway address details retrieved successfully.",
  "data": {
    "tenantDeviceId": "TENANT_DEVICE_OPAQUE_ID",
    "serverUrl": "https://axionpro-api.onrender.com",
    "serverPath": "/device-gateway",
    "serverPort": 443,
    "hasActiveGatewayUrl": true,
    "isReplacementPending": false,
    "replacementExpiresDateTime": null
  },
  "errors": []
}''',
    "Read only. Reads `TenantDeviceConfiguration`; never returns the ingress token or full working device URL.")

api(
    "Replace HTTPS gateway URL remotely",
    "POST /api/TenantDeviceConfiguration/replace-https-gateway-url",
    "Generates a pending opaque HTTPS gateway URL and queues a protected setting command while preserving the current URL until the device confirms the replacement.",
    "Tenant admin selects Replace URL for an already connected HTTPS device. Do not call it for heartbeat changes.",
    '''{
  "tenantDeviceId": "TENANT_DEVICE_OPAQUE_ID",
  "currentWebServerPassword": "CURRENT_DEVICE_WEB_PASSWORD",
  "replacementLifetimeMinutes": 30,
  "moduleId": 49,
  "operationId": 3
}''',
    SUCCESS_QUEUE.replace("Settings have been queued for the device.", "Device gateway URL replacement has been queued. The existing URL remains active until the device confirms the replacement."),
    "Updates `TenantDeviceConfiguration.PendingHttpsIngressTokenHash`, `PendingHttpsIngressTokenExpiresDateTime`, and audit fields. Inserts one protected `DeviceCommand` (`setdevinfo`). When device confirmation arrives, gateway processing promotes the pending token; `DeviceCommandResponse` stores acknowledgement.",
    ["The new raw URL is included only inside the protected command payload; it is not returned to Angular.", "If a pending replacement exists, wait for completion or expiry rather than submitting another one."])

api(
    "Apply base runtime configuration",
    "POST /api/TenantDeviceConfiguration/apply-runtime-configuration",
    "Queues heartbeat, optional volume and base managed device configuration. It can create an HTTPS gateway for a configured device that does not yet have one.",
    "Tenant admin saves the initial runtime configuration after device assignment. Use typed settings endpoints for detailed device sections.",
    '''{
  "tenantDeviceId": "TENANT_DEVICE_OPAQUE_ID",
  "currentWebServerPassword": "CURRENT_DEVICE_WEB_PASSWORD",
  "heartbeatIntervalSeconds": 20,
  "volume": 8,
  "disableLocalWebServer": true,
  "newWebServerPassword": null,
  "rebootAfterApply": true,
  "moduleId": 49,
  "operationId": 3
}''',
    '''{
  "isSucceeded": true,
  "message": "Tenant device configuration has been queued securely.",
  "data": {
    "configurationCommandId": 971,
    "configurationTrackingId": "b90a8374-42b2-42fc-bfad-88b590b6b811",
    "rebootCommandId": 972,
    "rebootTrackingId": "694f37f6-bc8d-4cad-8e21-8e3a36d5d1d0",
    "status": "Queued"
  },
  "errors": []
}''',
    "Updates `TenantDeviceConfiguration.HeartbeatIntervalSeconds`, `Configuration`, HTTPS token fields when missing, and audit fields. Inserts one protected `DeviceCommand` (`setdevinfo`) and optionally a second `DeviceCommand` (`reboot`).",
    ["`disableLocalWebServer` and `newWebServerPassword` are retained only for request compatibility. For local Web UI/API enablement or password rotation, use `settings/web-access`."])

api(
    "Queue device reboot",
    "POST /api/TenantDeviceConfiguration/reboot",
    "Queues a reboot through the configured outbound transport; no direct request is made to the device LAN address.",
    "Tenant admin explicitly presses Reboot after a confirmed setting change.",
    '''{
  "tenantDeviceId": "TENANT_DEVICE_OPAQUE_ID",
  "moduleId": 49,
  "operationId": 3
}''',
    SUCCESS_QUEUE.replace("Settings have been queued for the device.", "Tenant device reboot has been queued securely."),
    "Inserts one `DeviceCommand` (`reboot`). No `TenantDeviceConfiguration` setting is overwritten.")

api(
    "Move assigned device to another Tenant location",
    "POST /api/TenantDevice/update-location",
    "Moves the Tenant-owned installed device to a valid location in the same Tenant. It does not expose or alter the Host-owned DeviceMaster assignment.",
    "Tenant admin updates the physical location of an already assigned device.",
    '''{
  "tenantDeviceId": "TENANT_DEVICE_OPAQUE_ID",
  "tenantLocationId": 12,
  "moduleId": 49,
  "operationId": 3
}''',
    '''{
  "isSucceeded": true,
  "message": "Device location has been updated.",
  "data": {
    "id": "TENANT_DEVICE_OPAQUE_ID",
    "tenantLocationId": 12,
    "tenantLocationName": "Mumbai Office",
    "deviceCode": "MUM-GATE-01",
    "isActive": true
  },
  "errors": []
}''',
    "Updates `TenantDevice.TenantLocationId`, `UpdatedById`, and `UpdatedDateTime`. This endpoint does not create a device command.",
    ["Use the returned device's current location for the next employee enrollment eligibility check."])

api(
    "Publish the next MQTTS queue item now",
    "POST /api/TenantDeviceConfiguration/dispatch-mqtts-now",
    "Requests immediate broker delivery of the next already queued command for an MQTTS device only.",
    "Tenant admin presses Send now after a command is queued for an MQTTS device. Do not show this for HTTPS devices.",
    '''{
  "tenantDeviceId": "TENANT_DEVICE_OPAQUE_ID",
  "moduleId": 49,
  "operationId": 3
}''',
    '''{
  "isSucceeded": true,
  "message": "The next queued command has been published through secure MQTT.",
  "data": {"wasDispatched": true, "message": "The next queued command has been published through secure MQTT."},
  "errors": []
}''',
    "Does not insert a new business record. It reads and updates the next eligible `DeviceCommand` to publishing/awaiting-response state; device acknowledgement inserts `DeviceCommandResponse`.")

heading("2 Typed device settings", 1)
para("All endpoints in this section queue exactly one protected `setdevinfo` command into `DeviceCommand`. They validate the Tenant device, permission scope and current local web password first. They do not update a device directly and they do not update the `TenantDeviceConfiguration.Configuration` JSON. Use the matching dropdown API before rendering each section.")

setting_common = '''All calls include `tenantDeviceId`, `currentWebServerPassword`, `moduleId`, and `operationId`. All return the same queued-command response shown below.'''
para(setting_common)
code_block(SUCCESS_QUEUE)

settings = [
    ("Time settings", "POST /api/TenantDeviceConfiguration/settings/time", "GET /api/device-ddl-options/time", '''{
  "tenantDeviceId":"TENANT_DEVICE_OPAQUE_ID", "currentWebServerPassword":"CURRENT_DEVICE_WEB_PASSWORD", "timeFormat":24, "dateFormat":1,
  "daylightSavingEnabled":false, "daylightSavingStart":"3/21", "daylightSavingEnd":"9/21", "networkTimeEnabled":true,
  "timeZone":5, "rebootTime1":"00:00", "rebootTime2":"00:00", "rebootTime3":"00:00", "moduleId":49, "operationId":3
}''', "Time/date display, DST, NTP and scheduled reboot windows."),
    ("Synchronize device clock", "POST /api/TenantDeviceConfiguration/settings/time/sync", "GET /api/device-ddl-options/time", '''{
  "tenantDeviceId":"TENANT_DEVICE_OPAQUE_ID", "currentWebServerPassword":"CURRENT_DEVICE_WEB_PASSWORD", "utcDateTime":"2026-09-08T10:30:00Z", "moduleId":49, "operationId":3
}''', "Sets device clock. `utcDateTime` may be omitted to use current AxionPro UTC time."),
    ("Bell settings", "POST /api/TenantDeviceConfiguration/settings/bell", "GET /api/device-ddl-options/bell", '''{
  "tenantDeviceId":"TENANT_DEVICE_OPAQUE_ID", "currentWebServerPassword":"CURRENT_DEVICE_WEB_PASSWORD", "bellCount":0, "ringStyle":1, "bellOutput":0, "moduleId":49, "operationId":3
}''', "Bell count, ring pattern and output state."),
    ("Device setup", "POST /api/TenantDeviceConfiguration/settings/device-setup", "GET /api/device-ddl-options/device-setup", '''{
  "tenantDeviceId":"TENANT_DEVICE_OPAQUE_ID", "currentWebServerPassword":"CURRENT_DEVICE_WEB_PASSWORD", "language":1, "voiceVolume":8,
  "announcePersonName":true, "detectMultipleFaces":true, "resultDisplaySeconds":3, "screenSaverIdleSeconds":60, "sleepModeSeconds":15,
  "screenWakeUpMethod":1, "faceWakeUpSeconds":0, "resultDisplayStyle":1, "faceRecognitionDistance":3,
  "livenessDetectionEnabled":true, "showAvatar":true, "moduleId":49, "operationId":3
}''', "Screen, voice, wake-up and recognition experience."),
    ("Advanced settings", "POST /api/TenantDeviceConfiguration/settings/advanced", "GET /api/device-ddl-options/advanced", '''{
  "tenantDeviceId":"TENANT_DEVICE_OPAQUE_ID", "currentWebServerPassword":"CURRENT_DEVICE_WEB_PASSWORD", "maximumAdministrators":10,
  "verificationMode":1, "qrCodeMode":0, "hidePrivacyInformation":false, "faceMatchThreshold":50, "livenessThreshold":50,
  "fingerprintMatchThreshold":50, "fingerprintsPerUser":2, "maskDetectionEnabled":false, "maskThreshold":50, "fillLightMode":0,
  "constantFillLightPeriod":"00:00~00:00", "exposureCompensation":0, "palmVeinMatchThreshold":50, "palmDetectionThreshold":50,
  "disableFaceRecognition":false, "onlineDebugEnabled":false, "moduleId":49, "operationId":3
}''', "Verification, QR, privacy, recognition thresholds and fill-light options."),
    ("Door and lock settings", "POST /api/TenantDeviceConfiguration/settings/lock", "GET /api/device-ddl-options/lock", '''{
  "tenantDeviceId":"TENANT_DEVICE_OPAQUE_ID", "currentWebServerPassword":"CURRENT_DEVICE_WEB_PASSWORD", "doorOpenDelaySeconds":5,
  "doorSensorMode":1, "doorSensorDelaySeconds":10, "blockStrangerAccess":false, "doorPassword":0, "requiredUsersForDoorOpen":0,
  "antiPassbackMode":0, "wiegandOutput":0, "wiegandFormat":34, "accessLimit":0, "cardDisplayFormat":0,
  "reverseCardPin":false, "reverseWiegandOutput":false, "externalWiegandSnapshotEnabled":false, "interlockEnabled":false,
  "alarmProcessingEnabled":false, "failedVerificationLimit":0, "timeZonePunchLimit":0, "suppressAccessDeniedLog":false,
  "denyOutsideNormallyOpenTimeZone":false, "moduleId":49, "operationId":3
}''', "Door sensor, Wiegand, card display and access-control behavior."),
    ("Serial settings", "POST /api/TenantDeviceConfiguration/settings/serial", "GET /api/device-ddl-options/serial", '''{
  "tenantDeviceId":"TENANT_DEVICE_OPAQUE_ID", "currentWebServerPassword":"CURRENT_DEVICE_WEB_PASSWORD", "deviceAddress":1,
  "networkPort":5005, "baudRate":38400, "serialFunction":0, "moduleId":49, "operationId":3
}''', "Serial address, network port, baud rate and serial function."),
    ("Ethernet settings", "POST /api/TenantDeviceConfiguration/settings/ethernet", "GET /api/device-ddl-options/ethernet", '''{
  "tenantDeviceId":"TENANT_DEVICE_OPAQUE_ID", "currentWebServerPassword":"CURRENT_DEVICE_WEB_PASSWORD", "dhcpEnabled":false,
  "ipAddress":"192.168.1.224", "subnetMask":"255.255.255.0", "gateway":"192.168.1.1", "dnsServer":"8.8.8.8",
  "hideIpAddress":false, "moduleId":49, "operationId":3
}''', "Ethernet DHCP or static address. Static values are applied only when DHCP is false."),
    ("Wi Fi settings", "POST /api/TenantDeviceConfiguration/settings/wifi", "GET /api/device-ddl-options/wifi", '''{
  "tenantDeviceId":"TENANT_DEVICE_OPAQUE_ID", "currentWebServerPassword":"CURRENT_DEVICE_WEB_PASSWORD", "dhcpEnabled":false,
  "ipAddress":"192.168.1.225", "subnetMask":"255.255.255.0", "gateway":"192.168.1.1", "moduleId":49, "operationId":3
}''', "Wi-Fi DHCP or static address. Static values are applied only when DHCP is false."),
    ("App notification", "POST /api/TenantDeviceConfiguration/settings/app-notification", "GET /api/device-ddl-options/app-notification", '''{
  "tenantDeviceId":"TENANT_DEVICE_OPAQUE_ID", "currentWebServerPassword":"CURRENT_DEVICE_WEB_PASSWORD", "appNotificationEnabled":true,
  "appToken":"THIRD_PARTY_APP_TOKEN", "notificationType":1, "moduleId":49, "operationId":3
}''', "Optional third-party mobile/app notification. App token is protected and never returned."),
    ("Local web access and password", "POST /api/TenantDeviceConfiguration/settings/web-access", "No dropdown", '''{
  "tenantDeviceId":"TENANT_DEVICE_OPAQUE_ID", "currentWebServerPassword":"CURRENT_DEVICE_WEB_PASSWORD", "localWebServerEnabled":false,
  "newWebServerPassword":"NEW_DEVICE_WEB_PASSWORD", "moduleId":49, "operationId":3
}''', "Turns the device local Web UI/API on/off and, when supplied, rotates its password."),
    ("Physical screen menu PIN", "POST /api/TenantDeviceConfiguration/settings/screen-menu-pin", "No dropdown", '''{
  "tenantDeviceId":"TENANT_DEVICE_OPAQUE_ID", "currentWebServerPassword":"CURRENT_DEVICE_WEB_PASSWORD", "screenMenuPin":"1234", "moduleId":49, "operationId":3
}''', "Locks the device physical System/Local Manager screen menu behind a numeric PIN."),
]
for title, path, ddl, body, description in settings:
    api(title, path, description, f"Load `{ddl}` before this settings form when it is a dropdown section; submit only when the user presses Save for this section.", body, SUCCESS_QUEUE, "Inserts one protected `DeviceCommand` row. The settings are applied only when the device receives and acknowledges it.", ["Do not combine unrelated sections in one request. Refresh device command status after submit."])

heading("3 Dropdown APIs", 1)
para("Dropdown API calls are read-only. Each is called when its matching UI section opens, not globally at application startup. The response has `data` as a list of fields; each field has a stable `key` and `options` with `value` and `label`.")
DDL_RESPONSE = '''{
  "isSucceeded": true,
  "message": "Success",
  "data": [{
    "key": "credentialType",
    "label": "Credential type",
    "options": [{"value":"1","label":"Face biometric"}, {"value":"2","label":"Access card"}, {"value":"3","label":"Device PIN"}]
  }],
  "errors": []
}'''
table(["Endpoint", "Call when this section opens", "Primary returned fields", "Database effect"], [
    ["GET /api/device-ddl-options/time", "Time settings", "timeFormat, dateFormat, timeZone", "None"],
    ["GET /api/device-ddl-options/bell", "Bell settings", "ringStyle, bellOutput", "None"],
    ["GET /api/device-ddl-options/device-setup", "Device setup", "language, screenWakeUpMethod, resultDisplayStyle, faceRecognitionDistance", "None"],
    ["GET /api/device-ddl-options/advanced", "Advanced settings", "verificationMode, qrCodeMode, fillLightMode", "None"],
    ["GET /api/device-ddl-options/lock", "Lock settings", "doorSensorMode, antiPassbackMode, wiegandOutput, wiegandFormat, cardDisplayFormat", "None"],
    ["GET /api/device-ddl-options/serial", "Serial settings", "baudRate, serialFunction", "None"],
    ["GET /api/device-ddl-options/ethernet", "Ethernet settings", "DHCP options", "None"],
    ["GET /api/device-ddl-options/wifi", "Wi-Fi settings", "DHCP options", "None"],
    ["GET /api/device-ddl-options/app-notification", "App notification", "notificationType", "None"],
    ["GET /api/device-ddl-options/employee-device-credentials", "Remove credential selector", "credentialType 1 Face, 2 Card, 3 PIN", "None"],
    ["GET /api/device-ddl-options/employee-device-access-windows", "Employee device schedule", "dayOfWeek 1 Monday through 7 Sunday", "None"],
    ["GET /api/device-ddl-options/tenant-card-inventory", "Host card form/filter", "taxTreatment, cardStatus", "None"],
], [2.55, 1.5, 2.25, 0.8], 7.4)
para("Response example used by every dropdown endpoint", bold_prefix="Response example used by every dropdown endpoint")
code_block(DDL_RESPONSE)

heading("4 Host card inventory", 1)
doc.add_picture(str(CARD_FLOW), width=Inches(6.85))
para("Card inventory is Host-only. It is Tenant-scoped through the Host-selected opaque `tenantId`. A card number is encrypted in storage and the UI receives only `maskedCardNumber`. A card must be active and `Available` before a Tenant can bind it to an employee device enrollment.")

CARD_BODY = '''{
  "tenantId":"TENANT_OPAQUE_ID", "cardNumber":"6913186", "cardReference":"DEL-2026-0007", "purchaseCurrencyCode":"INR",
  "unitPurchasePriceExcludingTax":100.00, "supplierCountryId":101, "supplierStateId":23, "placeOfSupplyCountryId":101,
  "placeOfSupplyStateId":23, "supplierName":"Card Supplier", "supplierTaxRegistrationNumber":"23ABCDE1234F1Z5",
  "purchaseInvoiceNumber":"INV-2026-88", "purchaseInvoiceDate":"2026-09-08", "taxTreatment":1,
  "cgstRate":9, "cgstAmount":9, "sgstRate":9, "sgstAmount":9, "igstRate":0, "igstAmount":0,
  "foreignTaxLabel":null, "foreignTaxRate":0, "foreignTaxAmount":0, "customsDutyAmount":0, "freightAmount":5,
  "isActive":true, "moduleId":49, "operationId":1
}'''
CARD_RESPONSE = '''{
  "isSucceeded": true,
  "message": "Tenant card inventory item created successfully.",
  "data": {
    "id":"TENANT_CARD_OPAQUE_ID", "tenantId":"TENANT_OPAQUE_ID", "maskedCardNumber":"****3186", "cardReference":"DEL-2026-0007",
    "cardStatus":1, "purchaseCurrencyCode":"INR", "unitPurchasePriceExcludingTax":100.00, "cgstAmount":9.00,
    "sgstAmount":9.00, "igstAmount":0.00, "freightAmount":5.00, "landedCost":123.00, "isActive":true
  },
  "errors": []
}'''
api("Create card inventory item", "POST /api/TenantCardMaster/create", "Registers a physical card as available inventory for a selected Tenant.", "Host user receives cards from procurement and needs a card available for later employee assignment.", CARD_BODY, CARD_RESPONSE, "Inserts `TenantCardMaster`. Stores encrypted card number and a lookup hash; calculates `LandedCost` server side. No device command is inserted.")
api("Read one card", "GET /api/TenantCardMaster/get-by-id/{tenantCardId}?tenantId={tenantId}&moduleId={moduleId}&operationId={operationId}", "Loads an editable card detail without exposing its number.", "Host opens Card Inventory edit detail.", None, CARD_RESPONSE.replace("Tenant card inventory item created successfully.", "Success"), "Read only from `TenantCardMaster`.")
api("List card inventory", "GET /api/TenantCardMaster/get-all?tenantId={tenantId}&cardStatus=1&isActive=true&pageNumber=1&pageSize=10&moduleId={moduleId}&operationId={operationId}", "Lists inventory for Host card selection and filtering.", "Host opens inventory list or applies a status filter.", None, '''{
  "isSucceeded": true, "message":"Tenant card inventory retrieved successfully.",
  "data":[{"id":"TENANT_CARD_OPAQUE_ID","maskedCardNumber":"****3186","cardStatus":1,"isActive":true}],
  "pageNumber":1,"pageSize":10,"totalRecords":1,"totalPages":1,"errors":[]
}''', "Read only from `TenantCardMaster`.")
api("Update unassigned card", "POST /api/TenantCardMaster/update", "Corrects procurement, tax or reference information for a card that is not assigned.", "Host saves an unassigned card edit screen.", CARD_BODY.replace('"tenantId"', '"id":"TENANT_CARD_OPAQUE_ID", "tenantId"').replace('"operationId":1', '"operationId":3'), CARD_RESPONSE.replace("created", "updated"), "Updates `TenantCardMaster` audit fields. Assigned cards cannot be edited; unbind first.")
api("Set card active status", "POST /api/TenantCardMaster/update-status", "Activates or deactivates an unassigned card.", "Host uses the inventory status control.", '''{
  "id":"TENANT_CARD_OPAQUE_ID", "tenantId":"TENANT_OPAQUE_ID", "isActive":false, "moduleId":49, "operationId":3
}''', CARD_RESPONSE.replace('"isActive":true', '"isActive":false').replace("created", "status updated"), "Updates `TenantCardMaster.IsActive` and audit fields. An assigned card cannot be deactivated.")
api("Soft delete unassigned card", "DELETE /api/TenantCardMaster/delete/{tenantCardId}?tenantId={tenantId}&moduleId={moduleId}&operationId={operationId}", "Retires a card inventory record without leaking its number.", "Host removes an unassigned inventory card.", None, '''{"isSucceeded":true,"message":"Tenant card inventory item deleted successfully.","data":true,"errors":[]}''', "Sets `TenantCardMaster.IsActive = false`, `IsSoftDeleted = true`, soft-delete audit fields. Delete is rejected while assigned to an active employee enrollment.")

heading("5 Employee device enrollment and credentials", 1)
doc.add_picture(str(EMPLOYEE_FLOW), width=Inches(6.85))
para("The `EmployeeDeviceEnrollment` feature maps an existing AxionPro employee to an already configured Tenant device. The device-side `enrollid` is the backend-decrypted global Employee database ID. Angular never sees or submits it. Device privilege is forced to User (`admin = 0`) by the backend.")

api("Create employee device enrollment", "POST /api/EmployeeDeviceEnrollment/create", "Creates a tenant-safe employee-to-device mapping and queues the baseline device user.", "Tenant admin selects an employee and device after the employee has an active, attendance-enabled location assignment matching the device location.", '''{
  "employeeId":"EMPLOYEE_OPAQUE_ID", "tenantDeviceId":"TENANT_DEVICE_OPAQUE_ID", "tenantCardId":null,
  "accessEffectiveFromDateTime":"2026-09-09T00:00:00Z", "accessEffectiveToDateTime":null,
  "accessWindows":[
    {"dayOfWeek":1,"startLocalTime":"09:00:00","endLocalTime":"18:00:00","isActive":true},
    {"dayOfWeek":2,"startLocalTime":"09:00:00","endLocalTime":"18:00:00","isActive":true}
  ],
  "isActive":true, "moduleId":49, "operationId":1
}''', SUCCESS_ENROLLMENT, "Inserts `EmployeeDeviceEnrollment` and `EmployeeDeviceAccessWindow` rows. The backend queues protected baseline `setuserinfo` in `DeviceCommand`; vendor acknowledgement writes `DeviceCommandResponse`.", ["`tenantCardId` is not currently bound by this create handler. First create the enrollment, then call `card/bind`.", "The employee/device cannot later be swapped with update; create a new mapping instead."])

api("Read one enrollment", "GET /api/EmployeeDeviceEnrollment/get-by-id/{enrollmentId}?moduleId={moduleId}&operationId={operationId}", "Reads safe enrollment data and live device command statuses.", "Enrollment detail opens, or after any face/card/PIN/status action is submitted.", None, SUCCESS_ENROLLMENT.replace("Employee device enrollment created successfully.", "Success"), "Read only from `EmployeeDeviceEnrollment`, `EmployeeDeviceAccessWindow`, `TenantCardMaster` and referenced `DeviceCommand` rows. Raw face bytes, PIN, card number, Employee database ID and device serial are absent.")
api("List enrollments", "GET /api/EmployeeDeviceEnrollment/get-all?tenantDeviceId={id}&isActive=true&pageNumber=1&pageSize=10&moduleId={moduleId}&operationId={operationId}", "Lists device users for a Tenant device or employee.", "Device user list page and filtered search.", None, '''{
  "isSucceeded":true,"message":"Employee device enrollments retrieved successfully.",
  "data":[{"id":"ENROLLMENT_OPAQUE_ID","employeeName":"Ananya Sharma","deviceCode":"DELHI-GATE-01","isActive":true,"userActivationCommandStatus":4}],
  "pageNumber":1,"pageSize":10,"totalRecords":1,"totalPages":1,"errors":[]
}''', "Read only from enrollment tables and command references.")
api("Update access windows and validity", "POST /api/EmployeeDeviceEnrollment/update", "Changes employee-specific device schedule and validity dates only.", "Tenant admin changes working access time for the device user.", '''{
  "id":"ENROLLMENT_OPAQUE_ID", "employeeId":"EMPLOYEE_OPAQUE_ID", "tenantDeviceId":"TENANT_DEVICE_OPAQUE_ID",
  "accessEffectiveFromDateTime":"2026-09-09T00:00:00Z", "accessEffectiveToDateTime":null,
  "accessWindows":[{"dayOfWeek":1,"startLocalTime":"10:00:00","endLocalTime":"19:00:00","isActive":true}],
  "isActive":true, "moduleId":49, "operationId":3
}''', SUCCESS_ENROLLMENT.replace("created", "updated"), "Soft-deletes replaced child `EmployeeDeviceAccessWindow` rows and inserts new window rows. Updates `EmployeeDeviceEnrollment` validity/audit fields. This does not send a device command.", ["The submitted `isActive` must match current status. To enable/disable use `update-status`; a mismatched value is rejected so a device command cannot be skipped."])
api("Enable or disable device user", "POST /api/EmployeeDeviceEnrollment/update-status", "Changes local mapping status and queues vendor `enableuser` using the server-owned device enrollid.", "Tenant admin toggles employee availability on one device.", '''{
  "id":"ENROLLMENT_OPAQUE_ID", "isActive":false, "moduleId":49, "operationId":3
}''', SUCCESS_ENROLLMENT.replace('"userActivationCommandStatus": null', '"userActivationCommandStatus": 1').replace('"isActive": true', '"isActive": false').replace("created", "status updated"), "Updates `EmployeeDeviceEnrollment.IsActive`, audit fields and `UserActivationDeviceCommandId`. Inserts protected `DeviceCommand` payload `enableuser` with `enflag: false` or `true`.", ["False does not delete face, card or PIN. True re-enables the same device user and does not resend credentials.", "Do not consider device change complete until `userActivationCommandStatus` is `Completed` (4)."])
api("Upload employee face", "POST /api/EmployeeDeviceEnrollment/face/upsert", "Queues the selected Angular photo for device facial credential enrollment.", "Tenant admin selects a JPEG/PNG employee photo and presses Upload to device.", '''Content-Type: multipart/form-data
enrollmentId = ENROLLMENT_OPAQUE_ID
faceImage = <employee-face.jpg>   // image/jpeg or image/png, maximum 2 MB
moduleId = 49
operationId = 1''', SUCCESS_ENROLLMENT.replace('"faceDeploymentStatus": 0', '"faceDeploymentStatus": 1').replace('"faceCommandStatus": null', '"faceCommandStatus": 1').replace("created", "face enrollment has been queued for the device"), "Updates `EmployeeDeviceEnrollment.FaceImageHash`, `FaceDeploymentStatus`, `FaceDeviceCommandId` and audit fields. Inserts protected `DeviceCommand` (`setuserinfo`, backup 50). Raw image is not stored in any business table.")
api("Set or replace employee device PIN", "POST /api/EmployeeDeviceEnrollment/pin/upsert", "Queues a write-only 4 to 12 digit PIN for the employee on this device.", "Tenant admin enters PIN twice in UI, client validates equality, then submits once.", '''{
  "enrollmentId":"ENROLLMENT_OPAQUE_ID", "pin":"1234", "moduleId":49, "operationId":1
}''', SUCCESS_ENROLLMENT.replace('"pinDeploymentStatus": 0', '"pinDeploymentStatus": 1').replace('"pinCommandStatus": null', '"pinCommandStatus": 1').replace("created", "PIN enrollment has been queued for the device"), "Updates `EmployeeDeviceEnrollment.PinDeploymentStatus`, `PinDeviceCommandId` and audit fields. Inserts protected `DeviceCommand` (`setuserinfo`, backup 10). PIN is never stored or returned.")
api("Bind issued card to employee device user", "POST /api/EmployeeDeviceEnrollment/card/bind", "Assigns one active Available Host-issued card to this employee/device user and queues card enrollment.", "Tenant admin chooses an Available card from an inventory list that exposes masked numbers only.", '''{
  "enrollmentId":"ENROLLMENT_OPAQUE_ID", "tenantCardId":"TENANT_CARD_OPAQUE_ID", "moduleId":49, "operationId":1
}''', SUCCESS_ENROLLMENT.replace('"tenantCardId": null', '"tenantCardId":"TENANT_CARD_OPAQUE_ID"').replace('"maskedCardNumber": null', '"maskedCardNumber":"****3186"').replace('"cardDeploymentStatus": 0', '"cardDeploymentStatus": 1').replace('"cardCommandStatus": null', '"cardCommandStatus": 1').replace("created", "Card binding has been queued for the device"), "Updates `EmployeeDeviceEnrollment.TenantCardMasterId`, `CardDeploymentStatus`, `CardDeviceCommandId` and audit fields. Updates `TenantCardMaster.CardStatus` from Available (1) to Assigned (2). Inserts protected `DeviceCommand` (`setuserinfo`, backup 11).")
api("Remove one employee credential", "POST /api/EmployeeDeviceEnrollment/credential/remove", "Queues deletion of exactly Face, Card or PIN, without deleting the complete user.", "Tenant admin chooses Remove Face, Remove Card or Remove PIN. Load credential types from device-ddl-options first.", '''{
  "enrollmentId":"ENROLLMENT_OPAQUE_ID", "credentialType":2, "moduleId":49, "operationId":4
}''', SUCCESS_ENROLLMENT.replace('"tenantCardId": null', '"tenantCardId":null').replace("created", "Credential removal has been queued for the device"), "Inserts protected `DeviceCommand` (`deleteuser`) with backup number 50 Face, 11 Card, or 10 PIN. Updates enrollment credential state. Card removal also clears `TenantCardMasterId` and returns `TenantCardMaster.CardStatus` to Available.")
api("Remove complete employee device user", "DELETE /api/EmployeeDeviceEnrollment/delete/{enrollmentId}?moduleId={moduleId}&operationId={operationId}", "Removes the full device user and retires the mapping.", "Tenant admin confirms Remove employee from this device.", None, '''{"isSucceeded":true,"message":"Employee device enrollment deleted successfully.","data":true,"errors":[]}''', "Inserts protected `DeviceCommand` (`deleteuser`, full user backup 12), soft-deletes `EmployeeDeviceEnrollment`, and returns any attached `TenantCardMaster` to Available. Child schedule rows remain historic under the soft-deleted enrollment.")

heading("6 Employee location and work configuration", 1)
para("These endpoints support employee eligibility and hybrid/WFH configuration. They do not configure a biometric device by themselves. In particular, an active employee location assignment matching the device location is required before EmployeeDeviceEnrollment can be created.")

api("Create employee location assignment", "POST /api/EmployeeLocationAssignment/create", "Makes an employee eligible for a Tenant location and optionally makes it primary.", "Before adding an employee to a device at that location.", '''{
  "employeeId":"EMPLOYEE_OPAQUE_ID", "tenantLocationId":12, "isPrimary":true, "isAttendanceAllowed":true,
  "effectiveFrom":"2026-09-09", "effectiveTo":null, "isActive":true, "moduleId":44, "operationId":1
}''', '''{"isSucceeded":true,"message":"Employee location assignment created successfully.","data":{"id":81,"employeeId":501,"tenantLocationId":12,"tenantLocationName":"Delhi Office","isAttendanceAllowed":true,"isActive":true},"errors":[]}''', "Inserts `EmployeeLocationAssignment`. No DeviceCommand is created.")

WORK_CRUD = [
    ("Employee work arrangement", "EmployeeWorkArrangement", "EmployeeWorkArrangement", '''{"employeeId":"EMPLOYEE_OPAQUE_ID","attendancePolicyId":7,"primaryTenantLocationId":12,"workMode":3,"hybridType":1,"minimumOfficeDaysPerWeek":3,"effectiveFrom":"2026-09-09","effectiveTo":null,"isActive":true,"moduleId":44,"operationId":1}''', "Defines Office, Remote, Hybrid or WFH attendance context for one employee."),
    ("Employee work pattern", "EmployeeWorkPattern", "EmployeeWorkPattern", '''{"employeeWorkArrangementId":31,"dayOfWeek":1,"workMode":1,"tenantLocationId":12,"isWorkingDay":true,"isActive":true,"moduleId":44,"operationId":1}''', "Defines a weekday-level variation under a work arrangement."),
    ("Temporary work mode override", "EmployeeWorkModeOverride", "EmployeeWorkModeOverrideRequest", '''{"employeeId":"EMPLOYEE_OPAQUE_ID","employeeWorkArrangementId":31,"requestedWorkMode":2,"fromDate":"2026-09-10","toDate":"2026-09-12","tenantLocationId":null,"reason":"Approved travel","isActive":true,"moduleId":44,"operationId":1}''', "Records a temporary Office, Remote or WFH variation. Approval fields remain server-owned."),
]
for label, controller, table_name, create_body, purpose in WORK_CRUD:
    heading(label, 3)
    table(["Operation", "Endpoint", "Request or response", "Table effect"], [
        ["Create", f"POST /api/{controller}/create", create_body, f"Inserts `{table_name}`."],
        ["Read one", f"GET /api/{controller}/get-by-id/{{id}}?moduleId={{moduleId}}&operationId={{operationId}}", "No body. Returns one record.", f"Reads `{table_name}`."],
        ["List", f"GET /api/{controller}/get-all?...&moduleId={{moduleId}}&operationId={{operationId}}", "Filters plus pageNumber/pageSize. Returns paged list.", f"Reads `{table_name}`."],
        ["Update", f"POST /api/{controller}/update", "Same as create body plus `id`.", f"Updates `{table_name}` and audit fields."],
        ["Status", f"POST /api/{controller}/update-status", "`{\"id\": 31, \"isActive\": false, \"moduleId\":44, \"operationId\":3}`", f"Updates IsActive and audit fields."],
        ["Delete", f"DELETE /api/{controller}/delete/{{id}}?moduleId={{moduleId}}&operationId={{operationId}}", "No body. Returns `{ data: true }`.", f"Soft-deletes `{table_name}`."],
    ], [0.72, 1.55, 3.25, 1.15], 7.0)
    para(f"Why it exists: {purpose}", bold_prefix="Why it exists: ")
    code_block('''{
  "isSucceeded": true,
  "message": "Operation completed successfully.",
  "data": {"id": 31, "isActive": true},
    "errors": []
}''')

heading("Attendance punch validation and multi location rule", 2)
para("This is a required attendance rule, but it is not implemented by the configuration or enrollment APIs in this guide. `EmployeeLocationAssignment` decides whether a person is eligible to be enrolled on a device at a location. It does not yet decide whether a received biometric punch is accepted as attendance.")
table(["Business rule", "Required production decision"], [
    ["First IN", "When the employee has no open attendance session, accept the device event and open a session with its source device, location and device timestamp."],
    ["Second IN before OUT", "Reject the attendance decision even if the second device is at another office. Preserve the raw device event and record `DuplicateOpenIn` with the active session location, device and time."],
    ["OUT", "Close the current open session. The recommended default is to require the OUT at the same location as the active IN; a future Tenant attendance policy may explicitly permit a cross-location OUT."],
    ["IN after valid OUT", "Accept a new session. This supports an employee visiting Head Office in the morning and a Client Site or Delhi office later in the day."],
], [2.05, 4.95], 8.2)
para("The rule must apply to every employee and every enrolled device. It must be evaluated atomically on the server so two nearly simultaneous device punches cannot create two open IN sessions.")
para("Current device status: the temporary TIMMY `sendlog` diagnostic endpoint writes logs only and returns an acknowledgement; it does not persist attendance or evaluate duplicate IN events. The normal HTTPS gateway currently handles device polling and command delivery, not a completed attendance decision workflow.")
bullets([
    "Do not show a device punch as final attendance success in Angular until the future server attendance decision is available.",
    "The raw event must be append-only and retained whether accepted or rejected. It needs device serial, TenantDevice, location, encrypted/decrypted employee resolution, vendor log index, device event time, receive time, IN or OUT, verification mode, response decision and rejection reason.",
    "A device normally posts `sendlog` after local recognition. Therefore the server can reject the attendance record immediately, but an instant error on the device screen requires vendor support for real-time online authorization before the device finalizes the punch. That capability must be confirmed against the exact device protocol; it is not assumed by this API guide.",
])

heading("7 Generic device command submission", 1)
api("Submit an approved vendor command", "POST /api/device-commands/submit", "Queues a command from the protocol-approved catalog. This is not the normal form API for tenant UI; use typed settings and credential endpoints first.", "Only a controlled Host diagnostic/admin tool needs a supported command that has no typed user-facing endpoint.", '''{
  "tenantId":"TENANT_OPAQUE_ID", "tenantDeviceId":4, "commandName":"getdevinfo", "payload":"{\\"cmd\\":\\"getdevinfo\\"}",
  "moduleId":49, "operationId":1
}''', SUCCESS_QUEUE, "Inserts one `DeviceCommand`. On device delivery/acknowledgement, state and `DeviceCommandResponse` update. This endpoint must not be used to bypass tenant screen permissions or typed validation.")

heading("Implementation sequence for Angular", 1)
table(["Situation", "UI calls in order", "Done condition"], [
    ["New physical HTTPS device", "Host issue-bootstrap-url -> technician pastes one-time URL -> device calls bootstrap gateway -> Host assigns device/creates configuration -> Tenant apply-runtime-configuration", "Device appears assigned and runtime command completes."],
    ["Tenant changes device setting", "Open section -> GET matching device-ddl-options -> POST matching settings endpoint -> refresh command status", "DeviceCommand status is Completed."],
    ["Tenant adds employee to device", "Ensure EmployeeLocationAssignment -> POST enrollment/create -> GET enrollment -> face/card/PIN endpoints as needed -> poll GET enrollment", "Baseline and requested credential command statuses complete."],
    ["Tenant disables employee", "POST enrollment/update-status isActive false -> GET enrollment", "userActivationCommandStatus is Completed. Credentials are retained."],
    ["Host provides card", "POST TenantCardMaster/create -> card has Available status -> Tenant POST card/bind -> refresh enrollment", "Card command completed and inventory status is Assigned."],
    ["Tenant removes card", "POST credential/remove credentialType 2 -> refresh enrollment and card list", "Command completed and card inventory returns Available."],
], [1.45, 3.55, 1.65], 7.7)

heading("Status values UI should display", 1)
table(["Value", "Meaning", "UI treatment"], [
    ["Queued", "Durably stored but not sent", "Show Pending. Do not claim device changed."],
    ["Publishing", "Sending to transport", "Show Sending."],
    ["AwaitingResponse", "Device received/published path but acknowledgement is pending", "Show Waiting for device."],
    ["RetryScheduled", "Temporary delivery/ack issue; server will retry", "Show Retrying; disable duplicate submit where appropriate."],
    ["Completed", "Vendor result matched as successful", "Show Device updated."],
    ["Failed", "No successful acknowledgement or a vendor failure", "Show Device not confirmed. Allow an explicit retry based on user action."],
], [1.35, 3.0, 2.3], 8.2)

heading("Tables written by the new APIs", 1)
table(["Table", "Written by", "Purpose"], [
    ["DeviceInitialProvisioning", "issue-bootstrap-url", "Short-lived hashed initial device URL lifecycle."],
    ["TenantDeviceConfiguration", "base runtime, gateway replacement", "Connection metadata, heartbeat, gateway token hashes and audit fields."],
    ["TenantDevice", "update-location", "Assigned Tenant device and its current location."],
    ["DeviceCommand", "all remote setting, credential, reboot and generic command APIs", "Durable per-device command queue. Sensitive payloads are protected."],
    ["DeviceCommandResponse", "device inbound acknowledgement processing", "Parsed result and error record for a queued command."],
    ["DeviceMessageLog", "transport/gateway logging", "Inbound/outbound diagnostics with payload governance."],
    ["TenantCardMaster", "Host card inventory and card bind/remove", "Encrypted physical card inventory, lifecycle and procurement cost/tax details."],
    ["EmployeeDeviceEnrollment", "employee enrollment and credentials", "Employee/device mapping, status, command references and credential state."],
    ["EmployeeDeviceAccessWindow", "employee enrollment create/update", "Per-employee, per-device working access windows."],
    ["EmployeeLocationAssignment", "employee location APIs", "Location eligibility prerequisite for device enrollment."],
    ["EmployeeWorkArrangement", "work arrangement APIs", "Employee work arrangement, policy and effective period."],
    ["EmployeeWorkPattern", "work pattern APIs", "Weekly work arrangement pattern."],
    ["EmployeeWorkModeOverrideRequest", "override APIs", "Temporary requested work-mode deviation and approval lifecycle."],
], [1.95, 2.15, 2.55], 7.9)

para("End of guide. Use the response field names and endpoint paths as the Angular integration contract. Do not derive device IDs, device enrol IDs, raw card values, PINs, passwords or gateway tokens from UI state.")

doc.core_properties.title = "AxionPro Device UI API Developer Guide"
doc.core_properties.subject = "Device configuration and employee credential API handoff for Angular UI"
doc.core_properties.author = "Quecksilber Technologies"
doc.save(OUTPUT)
print(OUTPUT)
